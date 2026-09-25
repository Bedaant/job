"""DOCX resume generation from the Facts KB only (ADR-009: never the raw
resume blob). SPEC.md §3.5's 8 ATS-safety rules are satisfied by construction
here — documents/ats_safety.py::lint_docx re-checks the output independently
rather than trusting this file never regresses.
"""
import io

import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt

# No "Certification" heading exists in SPEC §3.5 rule 5's allowed set —
# certification facts fold into Education, the closest fit, rather than
# inventing a 6th heading the rule doesn't permit.
CATEGORY_TO_HEADING = {
    "experience": "Experience",
    "project": "Projects",
    "skill": "Skills",
    "certification": "Education",
    "education": "Education",
}
HEADING_ORDER = ["Experience", "Projects", "Skills", "Education"]
BODY_FONT = "Calibri"  # SPEC §3.5 rule 7's safe-font list, documents/ats_safety.py::SAFE_FONTS
DATE_RANGE_STYLE = "Job Copilot Date Range"


def _format_date_range(period_from, period_to) -> str | None:
    if period_from is None:
        return None
    end = period_to.strftime("%b %Y") if period_to else "Present"
    return f"{period_from.strftime('%b %Y')} – {end}"


def generate_resume_docx(headline: str | None, facts: list[dict]) -> bytes:
    document = docx.Document()

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    document.styles.add_style(DATE_RANGE_STYLE, WD_STYLE_TYPE.PARAGRAPH)

    if headline:
        document.add_paragraph("Summary", style="Heading 1")
        document.add_paragraph(headline)

    grouped: dict[str, list[dict]] = {}
    for fact in facts:
        heading = CATEGORY_TO_HEADING.get(fact["category"])
        if heading is None:
            continue
        grouped.setdefault(heading, []).append(fact)

    for heading in HEADING_ORDER:
        if heading not in grouped:
            continue
        document.add_paragraph(heading, style="Heading 1")
        for fact in grouped[heading]:
            document.add_paragraph(fact["achievement"], style="List Bullet")
            date_range = _format_date_range(fact.get("period_from"), fact.get("period_to"))
            if date_range:
                document.add_paragraph(date_range, style=DATE_RANGE_STYLE)

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
