"""SPEC.md §3.5 — the 8 ATS-safety rules, checked independently of
generate_docx.py's own construction (a linter that only trusted the generator
to stay correct wouldn't catch a future regression there).
"""
import io
import re

from docx import Document
from docx.oxml.ns import qn

DATE_RE = re.compile(r"^[A-Z][a-z]{2} \d{4} – ([A-Z][a-z]{2} \d{4}|Present)$")
ALLOWED_HEADINGS = {"Summary", "Experience", "Skills", "Education", "Projects"}
SAFE_FONTS = {"Calibri", "Arial", "Times New Roman", "Georgia"}
DATE_RANGE_STYLE = "Job Copilot Date Range"


def lint_docx(docx_bytes: bytes) -> list[str]:
    document = Document(io.BytesIO(docx_bytes))
    violations: list[str] = []

    # Rule 1: zero tables
    if document.tables:
        violations.append(f"Rule 1 (no tables) violated: {len(document.tables)} table(s) present")

    # Rule 2: zero text boxes
    if "txbxContent" in document.element.body.xml:
        violations.append("Rule 2 (no text boxes) violated: <w:txbxContent> present")

    # Rule 3: empty header and footer
    for section in document.sections:
        if "".join(p.text for p in section.header.paragraphs).strip():
            violations.append("Rule 3 (empty header/footer) violated: header is not empty")
        if "".join(p.text for p in section.footer.paragraphs).strip():
            violations.append("Rule 3 (empty header/footer) violated: footer is not empty")

    # Rule 4: zero embedded images
    if document.inline_shapes:
        violations.append(f"Rule 4 (no images) violated: {len(document.inline_shapes)} embedded image(s)")

    # Rule 5: headings only from the allowed set
    for p in document.paragraphs:
        if p.style.name.startswith("Heading") and p.text and p.text not in ALLOWED_HEADINGS:
            violations.append(f"Rule 5 (allowed headings only) violated: heading '{p.text}'")

    # Rule 6: single section, single column
    if len(document.sections) != 1:
        violations.append(f"Rule 6 (single section/column) violated: {len(document.sections)} sections")
    else:
        cols = document.sections[0]._sectPr.find(qn("w:cols"))
        if cols is not None and cols.get(qn("w:num")) not in (None, "1"):
            violations.append("Rule 6 (single section/column) violated: multi-column layout")

    # Rule 7: fonts limited to the safe list
    unsafe_fonts = {
        run.font.name
        for p in document.paragraphs
        for run in p.runs
        if run.font.name and run.font.name not in SAFE_FONTS
    }
    for font_name in sorted(unsafe_fonts):
        violations.append(f"Rule 7 (safe fonts only) violated: unsafe font '{font_name}'")

    # Rule 8: date ranges match the exact SPEC format
    for p in document.paragraphs:
        if p.style.name == DATE_RANGE_STYLE and p.text and not DATE_RE.match(p.text):
            violations.append(f"Rule 8 (date format) violated: '{p.text}'")

    return violations
