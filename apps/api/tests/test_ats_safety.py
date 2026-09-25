"""SPEC.md §3.5 — the 8 ATS-safety rules, asserted by inspecting the generated
DOCX XML directly (referenced by SPEC as already existing; it didn't, this is
that file). Plus the §3.5 parse-back check (fuzzy-match rendered bullets
against their source facts at >=0.9 similarity).
"""
import datetime

import docx

from documents.ats_safety import lint_docx
from documents.generate_docx import generate_resume_docx
from documents.parse_back import parse_back_check


def _facts():
    return [
        {
            "id": "f1", "category": "experience", "achievement": "Led a team of 5 engineers",
            "period_from": datetime.date(2022, 3, 1), "period_to": None,
        },
        {
            "id": "f2", "category": "project", "achievement": "Built a real-time matching pipeline",
            "period_from": datetime.date(2021, 6, 1), "period_to": datetime.date(2022, 1, 1),
        },
        {"id": "f3", "category": "skill", "achievement": "Python", "period_from": None, "period_to": None},
        {"id": "f4", "category": "certification", "achievement": "AWS Certified Solutions Architect",
         "period_from": None, "period_to": None},
    ]


# ---------- generate_resume_docx ----------

def test_generate_resume_docx_returns_valid_docx_bytes():
    import io
    docx_bytes = generate_resume_docx("Backend Engineer", _facts())
    document = docx.Document(io.BytesIO(docx_bytes))
    assert len(document.paragraphs) > 0


def test_generate_resume_docx_groups_facts_into_allowed_headings():
    import io
    docx_bytes = generate_resume_docx("Backend Engineer", _facts())
    document = docx.Document(io.BytesIO(docx_bytes))
    headings = {p.text for p in document.paragraphs if p.style.name.startswith("Heading")}
    assert headings <= {"Summary", "Experience", "Skills", "Education", "Projects"}
    assert "Experience" in headings  # from category=experience
    assert "Projects" in headings    # from category=project
    assert "Skills" in headings      # from category=skill
    assert "Education" in headings   # from category=certification, folded in (no "Certification" heading allowed)


def test_generate_resume_docx_renders_date_ranges_in_spec_format():
    import io
    docx_bytes = generate_resume_docx(None, _facts())
    document = docx.Document(io.BytesIO(docx_bytes))
    date_lines = [p.text for p in document.paragraphs if p.style.name == "Job Copilot Date Range"]
    assert "Mar 2022 – Present" in date_lines
    assert "Jun 2021 – Jan 2022" in date_lines


def test_generate_resume_docx_omits_summary_heading_when_no_headline():
    import io
    docx_bytes = generate_resume_docx(None, _facts())
    document = docx.Document(io.BytesIO(docx_bytes))
    headings = {p.text for p in document.paragraphs if p.style.name.startswith("Heading")}
    assert "Summary" not in headings


# ---------- lint_docx (the 8 rules) ----------

def test_lint_docx_passes_on_a_generator_produced_document():
    docx_bytes = generate_resume_docx("Backend Engineer", _facts())
    assert lint_docx(docx_bytes) == []


def test_lint_docx_flags_a_table():
    import io
    document = docx.Document()
    document.add_table(rows=1, cols=2)
    buf = io.BytesIO()
    document.save(buf)
    violations = lint_docx(buf.getvalue())
    assert any("table" in v.lower() for v in violations)


def test_lint_docx_flags_an_embedded_image():
    import io
    from PIL import Image

    document = docx.Document()
    img_buf = io.BytesIO()
    Image.new("RGB", (10, 10), color="white").save(img_buf, format="PNG")
    img_buf.seek(0)
    document.add_picture(img_buf)
    buf = io.BytesIO()
    document.save(buf)
    violations = lint_docx(buf.getvalue())
    assert any("image" in v.lower() for v in violations)


def test_lint_docx_flags_a_disallowed_heading():
    import io
    document = docx.Document()
    document.add_paragraph("Hobbies", style="Heading 1")
    buf = io.BytesIO()
    document.save(buf)
    violations = lint_docx(buf.getvalue())
    assert any("heading" in v.lower() for v in violations)


def test_lint_docx_flags_a_malformed_date_range():
    import io
    document = docx.Document()
    document.styles.add_style("Job Copilot Date Range", docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("March 2022 - Present", style="Job Copilot Date Range")  # wrong format
    buf = io.BytesIO()
    document.save(buf)
    violations = lint_docx(buf.getvalue())
    assert any("date" in v.lower() for v in violations)


def test_lint_docx_flags_an_unsafe_font():
    import io
    document = docx.Document()
    p = document.add_paragraph("Some text")
    p.runs[0].font.name = "Comic Sans MS"
    buf = io.BytesIO()
    document.save(buf)
    violations = lint_docx(buf.getvalue())
    assert any("font" in v.lower() for v in violations)


# ---------- parse_back_check ----------

def test_parse_back_check_passes_when_every_bullet_survives_the_round_trip():
    facts = _facts()
    docx_bytes = generate_resume_docx("Backend Engineer", facts)
    results = parse_back_check(docx_bytes, facts)
    assert all(r["passed"] for r in results)
    assert all(r["similarity"] >= 0.9 for r in results)


def test_parse_back_check_fails_a_fact_that_never_made_it_into_the_document():
    facts = _facts()
    docx_bytes = generate_resume_docx("Backend Engineer", facts)
    missing_fact = {"id": "f5", "category": "skill", "achievement": "Completely unrelated made-up text",
                     "period_from": None, "period_to": None}
    results = parse_back_check(docx_bytes, facts + [missing_fact])
    failed = [r for r in results if r["fact_id"] == "f5"]
    assert len(failed) == 1
    assert failed[0]["passed"] is False
    assert failed[0]["similarity"] < 0.9
