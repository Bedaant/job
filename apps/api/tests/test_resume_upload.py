import io
from unittest.mock import patch

import docx
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import main
import models
from database import Base, get_db

MOCK_FACTS = [
    {
        "category": "experience",
        "achievement": "Cut checkout API p99 latency from 1.4s to 180ms",
        "proof": "Payments team",
        "metric": "87% reduction",
        "tags": ["backend"],
    }
]


@pytest.fixture()
def client():
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,  # one shared connection — SQLite :memory: is per-connection otherwise
    )
    Base.metadata.create_all(engine)
    TestSessionLocal = sessionmaker(bind=engine)

    def override_get_db():
        db = TestSessionLocal()
        try:
            yield db
        finally:
            db.close()

    main.app.dependency_overrides[get_db] = override_get_db
    yield TestClient(main.app), TestSessionLocal
    main.app.dependency_overrides.clear()
    engine.dispose()


def _signup_and_login(client, email):
    client.post("/auth/signup", json={"email": email, "password": "correct horse battery staple"})
    resp = client.post("/auth/login", data={"username": email, "password": "correct horse battery staple"})
    return resp.json()["access_token"]


def _create_profile(session_factory, email):
    db = session_factory()
    user = db.query(models.User).filter(models.User.email == email).first()
    profile = models.Profile(user_id=user.id, persona=models.Persona.developer)
    db.add(profile)
    db.commit()
    db.refresh(profile)
    profile_id = profile.id
    db.close()
    return profile_id


def _resume_docx_bytes():
    document = docx.Document()
    document.add_paragraph("Jane Doe")
    document.add_paragraph("Cut checkout API p99 latency from 1.4s to 180ms at Acme Corp")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_full_upload_review_confirm_flow(client):
    test_client, SessionLocal = client
    token = _signup_and_login(test_client, "resume-flow@example.com")
    profile_id = _create_profile(SessionLocal, "resume-flow@example.com")
    headers = {"Authorization": f"Bearer {token}"}

    with patch("main.extract_facts_from_text", return_value=MOCK_FACTS):
        resp = test_client.post(
            f"/profiles/{profile_id}/resume",
            headers=headers,
            files={"file": ("resume.docx", _resume_docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
    assert resp.status_code == 200
    body = resp.json()
    assert body["status"] == "ready"
    assert len(body["facts"]) == 1
    upload_id = body["upload_id"]

    # poll endpoint returns the same result
    resp = test_client.get(f"/profiles/{profile_id}/resume/{upload_id}", headers=headers)
    assert resp.status_code == 200
    assert resp.json()["status"] == "ready"

    # facts are NOT yet in resume_facts until explicitly confirmed
    resp = test_client.get(f"/resume-facts?profile_id={profile_id}", headers=headers)
    assert resp.json() == []

    # confirm via facts:bulk
    resp = test_client.post(
        f"/profiles/{profile_id}/facts:bulk",
        headers=headers,
        json={"facts": body["facts"]},
    )
    assert resp.status_code == 200
    confirmed = resp.json()
    assert len(confirmed) == 1
    assert confirmed[0]["achievement"] == MOCK_FACTS[0]["achievement"]

    # now they show up
    resp = test_client.get(f"/resume-facts?profile_id={profile_id}", headers=headers)
    assert len(resp.json()) == 1


def test_upload_rejects_oversized_file(client):
    test_client, SessionLocal = client
    token = _signup_and_login(test_client, "big-file@example.com")
    profile_id = _create_profile(SessionLocal, "big-file@example.com")
    headers = {"Authorization": f"Bearer {token}"}

    oversized = b"x" * (5 * 1024 * 1024 + 1)
    resp = test_client.post(
        f"/profiles/{profile_id}/resume",
        headers=headers,
        files={"file": ("resume.pdf", oversized, "application/pdf")},
    )
    assert resp.status_code == 413


def test_cross_tenant_resume_upload_is_404(client):
    test_client, SessionLocal = client
    _signup_and_login(test_client, "owner2@example.com")
    owner_profile_id = _create_profile(SessionLocal, "owner2@example.com")

    intruder_token = _signup_and_login(test_client, "intruder2@example.com")
    _create_profile(SessionLocal, "intruder2@example.com")

    resp = test_client.post(
        f"/profiles/{owner_profile_id}/resume",
        headers={"Authorization": f"Bearer {intruder_token}"},
        files={"file": ("resume.docx", _resume_docx_bytes(), "application/octet-stream")},
    )
    assert resp.status_code == 404


def test_facts_bulk_never_auto_persists_from_parsing_alone(client):
    """ADR-009: parsed output is never silently trusted."""
    test_client, SessionLocal = client
    token = _signup_and_login(test_client, "trust-check@example.com")
    profile_id = _create_profile(SessionLocal, "trust-check@example.com")
    headers = {"Authorization": f"Bearer {token}"}

    with patch("main.extract_facts_from_text", return_value=MOCK_FACTS):
        test_client.post(
            f"/profiles/{profile_id}/resume",
            headers=headers,
            files={"file": ("resume.docx", _resume_docx_bytes(), "application/octet-stream")},
        )

    resp = test_client.get(f"/resume-facts?profile_id={profile_id}", headers=headers)
    assert resp.json() == [], "facts must not be persisted until facts:bulk is explicitly called"


def test_a_failed_parse_never_shows_the_raw_provider_error(client):
    """Found by the a11y audit: a bad/missing ANTHROPIC_API_KEY showed users
    "Error code: 401 - {'type': 'error', ... 'API key is invalid.'}" and told them
    to try another file format. The raw exception is internal; the user gets a
    plain message that points at the by-hand path instead."""
    test_client, SessionLocal = client
    token = _signup_and_login(test_client, "resume-fail@example.com")
    profile_id = _create_profile(SessionLocal, "resume-fail@example.com")
    raw = "Error code: 401 - {'type': 'error', 'error': {'type': 'authentication_error', 'message': 'API key is invalid.'}}"

    with patch("main.extract_facts_from_text", side_effect=RuntimeError(raw)):
        resp = test_client.post(
            f"/profiles/{profile_id}/resume",
            headers={"Authorization": f"Bearer {token}"},
            files={"file": ("resume.docx", _resume_docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )

    body = resp.json()
    assert body["status"] == "failed"
    assert "API key" not in body["error"] and "401" not in body["error"] and "{" not in body["error"]
    assert "by hand" in body["error"]
